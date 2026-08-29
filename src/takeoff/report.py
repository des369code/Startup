"""Deliverable report writers: JSON, CSV, XLSX.

All three consume the same TakeoffResult. CSV/Excel column order is fixed
and exact (the CLI depends on it). qa_flags is sheet-level QA — identical
in every row.
"""
import csv
import io
import json
from dataclasses import asdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .aggregate import TradeRollup
from .models import TakeoffResult

COLUMNS = [
    "sheet_name", "class_name_en", "class_name_ms", "measure",
    "quantity", "unit", "source_ids", "confidence", "qa_flags",
]


def _row(result: TakeoffResult, m) -> list:
    return [
        result.sheet_name, m.class_name_en, m.class_name_ms or "", m.measure,
        m.quantity, m.unit, "; ".join(m.source_ids), m.confidence,
        "; ".join(result.qa),
    ]


def to_json(result: TakeoffResult) -> str:
    """Full dataclass dump — every field, nothing flattened."""
    return json.dumps(asdict(result), indent=2)


def to_csv(result: TakeoffResult) -> str:
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\n")
    w.writerow(COLUMNS)
    for m in result.measurements:
        w.writerow(_row(result, m))
    return out.getvalue()


def write_xlsx(result: TakeoffResult, out_path: str) -> str:
    """One sheet 'takeoff', bold frozen header, one row per measurement."""
    wb = Workbook()
    ws = wb.active
    ws.title = "takeoff"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for m in result.measurements:
        ws.append(_row(result, m))
    ws.freeze_panes = "A2"
    wb.save(out_path)
    return out_path


def write_combined_takeoff(results: list[TakeoffResult], out_path: str) -> str:
    """Sheet 'takeoff' of the multi-sheet workbook: every measurement row
    from every result, same 9 columns, sheet_name populated per row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "takeoff"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for result in results:
        for m in result.measurements:
            ws.append(_row(result, m))
    ws.freeze_panes = "A2"
    wb.save(out_path)
    return out_path


def add_rollup_sheet(xlsx_path: str, rollups: list[TradeRollup]) -> str:
    """Append a 'rollup' sheet to an existing xlsx: one row per rollup (bold
    header: class_name_en, measure, total, unit, reconciled), sorted by
    class_name_en. reconciled is always False in v0 — see aggregate.py."""
    wb = load_workbook(xlsx_path)
    ws = wb.create_sheet("rollup")
    ws.append(["class_name_en", "measure", "total", "unit", "reconciled"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in sorted(rollups, key=lambda r: r.class_name_en):
        ws.append([r.class_name_en, r.measure, r.total, r.unit, r.reconciled])
    wb.save(xlsx_path)
    return xlsx_path
