"""CLI end-to-end: the product's only human entry point, exercised through
main() with a real fixture sheet — extraction to annotated PDF and Excel."""

import json
import shutil
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

import takeoff.cli as cli_mod
from takeoff.cli import main
from tests.fixtures import make_synthetic_drawing


def test_cli_full_run(tmp_path, capsys):
    pdf, truth = make_synthetic_drawing(tmp_path)
    out = tmp_path / "o"
    rc = main(["run", pdf, "--prompt", "", "--scale", "100", "--out", str(out), "--mock"])
    assert rc == 0
    assert (out / f"{Path(pdf).stem}-takeoff.xlsx").exists()   # matches CLI output contract
    assert "Asphalt" in capsys.readouterr().out


def test_cli_corrupt_pdf_exits_1(tmp_path, capsys):
    bad = tmp_path / "garbage.pdf"
    bad.write_bytes(b"not a pdf at all")
    rc = main(["run", str(bad), "--out", str(tmp_path / "o"), "--mock"])
    assert rc == 1
    assert "cannot read PDF" in capsys.readouterr().err


def test_cli_annotated_pdf_has_overlay(tmp_path):
    # the annotated PDF is the QA surface — assert it was written AND the quantity
    # label text is actually drawn on it (search page text for "50.00")
    pdf, _ = make_synthetic_drawing(tmp_path)
    out = tmp_path / "o"; main(["run", pdf, "--scale", "100", "--out", str(out), "--mock"])
    ann = fitz.open(out / f"{Path(pdf).stem}-annotated.pdf")
    assert any("50.00" in page.get_text() for page in ann)


def test_cli_dry_run_no_files(tmp_path, capsys):
    # cost estimate only: geometry extraction, no semantics, no output files
    pdf, _ = make_synthetic_drawing(tmp_path)
    out = tmp_path / "dryout"
    rc = main(["run", pdf, "--dry-run", "--out", str(out)])
    assert rc == 0
    assert not out.exists()
    text = capsys.readouterr().out
    assert "dry run" in text or "estimate" in text


def test_cli_mock_banner(tmp_path, capsys):
    # the banner is the demo-misuse guard — grep-able on stdout
    pdf, _ = make_synthetic_drawing(tmp_path)
    rc = main(["run", pdf, "--out", str(tmp_path / "o"), "--mock"])
    assert rc == 0
    assert "MOCK SEMANTICS" in capsys.readouterr().out


def test_cli_wrong_file_type_exits_1(tmp_path, capsys):
    # pymupdf can open a markdown file as a document; extract_sheet must
    # reject anything that is not a PDF at the single chokepoint.
    bad = tmp_path / "notes.md"
    bad.write_text("# Title\n\nJust some notes, not a drawing.\n")
    rc = main(["run", str(bad), "--out", str(tmp_path / "o"), "--mock"])
    assert rc == 1
    assert "cannot read PDF" in capsys.readouterr().err


def test_cli_dir_mode(tmp_path, capsys):
    # the wedge demo: a folder of sheets -> per-sheet artifacts PLUS one
    # combined rollup xlsx/json with two copies of the same sheet summed
    d = tmp_path / "sitesheets"; d.mkdir()
    p1, _ = make_synthetic_drawing(d)
    shutil.copyfile(p1, str(d / "sheet2.pdf"))   # second stem: distinct sheet name
    out = tmp_path / "o"
    rc = main(["run", str(d), "--out", str(out), "--mock"])
    assert rc == 0
    assert (out / "sitesheets-rollup.xlsx").exists()
    assert (out / "sitesheets-rollup.json").exists()
    text = capsys.readouterr().out
    assert "Asphalt" in text
    assert "100.0" in text   # 50 + 50 across the two sheets


def test_cli_dir_empty_exits_1(tmp_path, capsys):
    d = tmp_path / "empty"; d.mkdir()
    assert main(["run", str(d), "--out", str(tmp_path / "o"), "--mock"]) == 1
    assert "no PDF files found" in capsys.readouterr().err


def test_cli_dir_partial_corrupt_pdf_exits_1(tmp_path, capsys):
    # one corrupt sheet must NOT silently cancel the dir: the good sheet still
    # measures (its files exist), the failure is loud on stderr, rc = 1
    d = tmp_path / "sheets"; d.mkdir()
    p1, _ = make_synthetic_drawing(d)
    (d / "x.pdf").write_text("this is not a pdf, just text")
    out = tmp_path / "o"
    rc = main(["run", str(d), "--out", str(out), "--mock"])
    assert rc == 1
    assert "cannot read PDF" in capsys.readouterr().err
    assert (out / f"{Path(p1).stem}-takeoff.xlsx").exists()


def test_cli_negative_scale_exits_nonzero(tmp_path, capsys):
    # --scale must reject non-positive values before any work: argparse's
    # type validator exits rc 2 and NOTHING is written. (--scale 0 used to
    # zero every quantity; --scale -5 squared the factor into a wrong-but-
    # plausible 0.13 instead of 50.00.)
    pdf, _ = make_synthetic_drawing(tmp_path)
    out = tmp_path / "o"
    for bad in ("0", "-5"):
        with pytest.raises(SystemExit) as e:
            main(["run", pdf, "--scale", bad, "--out", str(out), "--mock"])
        assert e.value.code != 0
    assert "scale" in capsys.readouterr().err
    assert not out.exists()


def test_cli_review_lines_persist_to_qa_json(tmp_path, monkeypatch, capsys):
    # REVIEW lines must be folded into result.qa BEFORE artifacts are written,
    # so qa.json AND the xlsx qa_flags column carry them (not just stdout).
    fake_line = ("ASPHALT: drawing says 41.00 m — engine measured 50.0 m — REVIEW")
    monkeypatch.setattr(cli_mod, "reconcile_runs",
                        lambda result, words, runs: [fake_line])
    pdf, _ = make_synthetic_drawing(tmp_path)
    out = tmp_path / "o"
    rc = main(["run", pdf, "--out", str(out), "--mock"])
    assert rc == 0
    qa = json.loads((out / f"{Path(pdf).stem}-qa.json").read_text())
    assert fake_line in qa["qa"]
    wb = load_workbook(out / f"{Path(pdf).stem}-takeoff.xlsx")
    assert any(fake_line in str(c.value)
               for row in wb["takeoff"].iter_rows() for c in row)
    assert fake_line in capsys.readouterr().out  # stdout keeps the line


def test_cli_dir_dry_run_refused_no_files(tmp_path, capsys):
    # --dry-run + a directory would silently skip the guard and burn a real
    # client's money; refuse it and write nothing
    d = tmp_path / "sheets"; d.mkdir()
    make_synthetic_drawing(d)
    out = tmp_path / "dryout"
    rc = main(["run", str(d), "--dry-run", "--out", str(out), "--mock"])
    assert rc == 1
    assert "dry run" in capsys.readouterr().err
    assert not out.exists()
